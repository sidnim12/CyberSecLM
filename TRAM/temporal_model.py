import openpyxl
from collections import Counter


DATASET_PATH = "../../Datasets/ChronoCTI-main/ChronoCTI-main/Datasets/Dataset/temporal_relation_dataset.xlsx"


def load_training_edges():
    wb = openpyxl.load_workbook(DATASET_PATH, read_only=True)
    ws = wb["Sheet1"]

    edges = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        report, t1, t2, relation, comment, mask = row[:6]

        if relation == "NEXT" and mask == "train":
            edges.append((t1, t2))

    wb.close()
    return edges


def load_eval_edges():
    wb = openpyxl.load_workbook(DATASET_PATH, read_only=True)
    ws = wb["Sheet1"]

    edges = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        report, t1, t2, relation, comment, mask = row[:6]

        if relation == "NEXT" and mask == "eval":
            edges.append((t1, t2))

    wb.close()
    return edges


def build_transition_counts(edges):
    return Counter(edges)


def build_transition_probabilities(edges):
    transitions = Counter(edges)
    outgoing = Counter(t1 for t1, t2 in edges)

    probabilities = {}

    for (t1, t2), count in transitions.items():
        probabilities[(t1, t2)] = count / outgoing[t1]

    return probabilities


def evaluate_top_k(eval_edges, probabilities, k):
    correct = 0
    evaluated = 0

    for t1, actual_t2 in eval_edges:
        candidates = {
            next_t2: probability
            for (source_t1, next_t2), probability in probabilities.items()
            if source_t1 == t1
        }

        if not candidates:
            continue

        evaluated += 1

        top_k = sorted(
            candidates,
            key=candidates.get,
            reverse=True
        )[:k]

        if actual_t2 in top_k:
            correct += 1

    return correct, evaluated


def evaluate_mrr(eval_edges, probabilities):
    reciprocal_ranks = []

    for t1, actual_t2 in eval_edges:
        candidates = {
            next_t2: probability
            for (source_t1, next_t2), probability in probabilities.items()
            if source_t1 == t1
        }

        if not candidates:
            continue

        ranked_candidates = sorted(
            candidates,
            key=candidates.get,
            reverse=True
        )

        if actual_t2 in ranked_candidates:
            rank = ranked_candidates.index(actual_t2) + 1
            reciprocal_ranks.append(1 / rank)

    if not reciprocal_ranks:
        return 0

    return sum(reciprocal_ranks) / len(reciprocal_ranks)


def build_global_frequency_baseline(edges):
    next_techniques = Counter(t2 for t1, t2 in edges)
    return next_techniques

def evaluate_global_baseline(eval_edges, global_frequency, k):
    correct = 0
    evaluated = 0

    ranked_techniques = [
        technique
        for technique, count in global_frequency.most_common()
    ]

    for t1, actual_t2 in eval_edges:
        evaluated += 1

        top_k = ranked_techniques[:k]

        if actual_t2 in top_k:
            correct += 1

    return correct, evaluated

def analyze_transition_matrix(probabilities):
    print("\nTop transition probabilities:")

    top_transitions = sorted(
        probabilities.items(),
        key=lambda x: x[1],
        reverse=True
    )[:20]

    for (t1, t2), probability in top_transitions:
        print(f"{t1} -> {t2}: {probability:.4f}")

def plot_top_transitions(probabilities, n=20):
    import matplotlib.pyplot as plt

    top_transitions = sorted(
        probabilities.items(),
        key=lambda x: x[1],
        reverse=True
    )[:n]

    labels = [
        f"{t1.split(':')[0]} → {t2.split(':')[0]}"
        for (t1, t2), _ in top_transitions
    ]

    values = [
        probability
        for _, probability in top_transitions
    ]

    plt.figure(figsize=(10, 7))

    plt.barh(labels[::-1], values[::-1])

    plt.xlabel("Transition Probability")
    plt.ylabel("ATT&CK Transition")
    plt.title("Top ATT&CK Technique Transition Probabilities")

    plt.tight_layout()
    plt.show()

def analyze_branching(probabilities):
    from collections import defaultdict

    destinations = defaultdict(set)

    for (t1, t2), probability in probabilities.items():
        destinations[t1].add(t2)

    branching_counts = {
        t1: len(next_techniques)
        for t1, next_techniques in destinations.items()
    }

    print("\nTransition branching analysis:")

    total_sources = len(branching_counts)
    deterministic = sum(
        1 for count in branching_counts.values()
        if count == 1
    )
    low_branching = sum(
        1 for count in branching_counts.values()
        if 2 <= count <= 5
    )
    high_branching = sum(
        1 for count in branching_counts.values()
        if count > 5
    )

    average_branching = (
        sum(branching_counts.values()) / total_sources
        if total_sources else 0
    )

    print(f"Source techniques: {total_sources}")
    print(f"Deterministic (1 next): {deterministic}")
    print(f"Low branching (2-5 next): {low_branching}")
    print(f"High branching (>5 next): {high_branching}")
    print(f"Average next techniques: {average_branching:.2f}")

    print("\nMost branching source techniques:")

    for technique, count in sorted(
        branching_counts.items(),
        key=lambda x: x[1],
        reverse=True
    )[:10]:
        print(f"{technique}: {count} possible next techniques")

def analyze_transition_entropy(probabilities):
    from collections import defaultdict
    import math

    transitions = defaultdict(list)

    for (t1, t2), probability in probabilities.items():
        transitions[t1].append(probability)

    entropies = {}

    for t1, probs in transitions.items():
        entropy = -sum(
            p * math.log2(p)
            for p in probs
            if p > 0
        )

        entropies[t1] = entropy

    print("\nTransition entropy analysis:")

    print("\nLowest-entropy source techniques:")

    for technique, entropy in sorted(
        entropies.items(),
        key=lambda x: x[1]
    )[:10]:
        print(f"{technique}: {entropy:.4f} bits")

    print("\nHighest-entropy source techniques:")

    for technique, entropy in sorted(
        entropies.items(),
        key=lambda x: x[1],
        reverse=True
    )[:10]:
        print(f"{technique}: {entropy:.4f} bits")


def plot_entropy_vs_branching(probabilities):
    from collections import defaultdict
    import math
    import matplotlib.pyplot as plt

    transitions = defaultdict(list)

    for (t1, t2), probability in probabilities.items():
        transitions[t1].append(probability)

    branching = []
    entropy = []

    for t1, probs in transitions.items():
        branching.append(len(probs))

        h = -sum(
            p * math.log2(p)
            for p in probs
            if p > 0
        )

        entropy.append(h)

    plt.figure(figsize=(10, 7))

    plt.scatter(branching, entropy)

    plt.xlabel("Number of Possible Next Techniques")
    plt.ylabel("Transition Entropy (bits)")
    plt.title("Transition Entropy vs Branching")

    plt.tight_layout()
    plt.show()

if __name__ == "__main__":
    edges = load_training_edges()
    eval_edges = load_eval_edges()

    probabilities = build_transition_probabilities(edges)
    analyze_branching(probabilities)
    analyze_transition_matrix(probabilities)
    analyze_transition_entropy(probabilities)
    plot_entropy_vs_branching(probabilities)
    plot_top_transitions(probabilities)
    global_frequency = build_global_frequency_baseline(edges)

    print("\nMarkov Model Evaluation:")

    for k in [1, 3, 5]:
        correct, evaluated = evaluate_top_k(
            eval_edges,
            probabilities,
            k
        )

        accuracy = correct / evaluated if evaluated else 0

        print(
            f"Top-{k} Accuracy: {accuracy:.4f} "
            f"({correct}/{evaluated})"
        )

    mrr = evaluate_mrr(eval_edges, probabilities)
    print(f"MRR: {mrr:.4f}")

    known_sources = {t1 for t1, _ in probabilities}

    known_source_edges = [
        edge for edge in eval_edges
        if edge[0] in known_sources
    ]

    print("\nGlobal-frequency baseline (same evaluation set):")

    for k in [1, 3, 5]:
        correct, evaluated = evaluate_global_baseline(
            known_source_edges,
            global_frequency,
            k
        )

        accuracy = correct / evaluated if evaluated else 0

        print(
            f"Top-{k} Accuracy: {accuracy:.4f} "
            f"({correct}/{evaluated})"
        )